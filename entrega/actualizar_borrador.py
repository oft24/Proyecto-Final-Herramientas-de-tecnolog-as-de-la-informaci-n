"""Actualiza únicamente respuestas verificadas en la plantilla académica.

Uso: python entrega/actualizar_borrador.py <sha-de-remediacion>
No inserta capturas ni marca como terminadas las etapas AWS pendientes.
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document


WORD = Path(__file__).with_name("Evidencias_EntregaFinal_en_progreso.docx")
REPO = "https://github.com/oft24/Proyecto-Final-Herramientas-de-tecnolog-as-de-la-informaci-n"


def replace_text(paragraph, value: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)


def main(sha: str) -> None:
    document = Document(WORD)
    answers = {
        32: "El parche permitía pedir el reenvío con un código de pedido sin verificar sesión ni propietario. En la prueba local, otro usuario y una persona anónima recibieron 202 en vez de 404 y 401.",
        33: "Es control de acceso roto a nivel de objeto (CWE-639). La prueba de autorización, no Bandit, detectó el defecto y detuvo el pipeline local. Falta repetirla en QA.",
        35: "Le asigné severidad media: con un código de pedido conocido se podía activar una confirmación ajena, incluso repetidamente. El endpoint no devolvía el contenido del pedido al solicitante.",
        36: "El identificador aleatorio dificulta adivinarlo, pero no sustituye autorización si el código se comparte o filtra. La explotación y el alcance en QA aún deben verificarse.",
        38: "No atribuyo esta falla a un falso positivo. La prueba ejercitó la ruta y mostró 202 indebido; el pipeline local bloqueó por dos pruebas de autorización fallidas.",
        39: "Bandit no detectó esa regla de negocio. Sus avisos medios permanecen en el reporte, pero ninguno explica el acceso a pedidos ajenos; la corrida QA todavía está pendiente.",
        42: "La versión vulnerable quedó como commit de prueba local y no se levantó como servicio público. Para QA indiqué ejecutarla solo como pipeline aislado, sin abrir puertos; falta confirmar esa contención en AWS.",
        43: "No borré ni alteré pedidos reales para demostrarla: las pruebas locales usaron dobles de RDS y notificaciones. La contención no reemplaza el arreglo de autorización.",
        45: "Exigí sesión y comparación de orders.user_id con el usuario antes de notificar. Un pedido ajeno devuelve 404 sin evento, y una petición anónima devuelve 401; el propietario conserva el reenvío.",
        46: "Persistí el código público del pedido en RDS y tomé el correo de la cuenta autenticada. Probé casos propios, ajenos, anónimos y fallas; el pipeline local permitió la versión corregida. Falta validarla en QA.",
        50: "Lo más difícil fue relacionar el código BDK visible con la fila real de RDS: antes no se persistía ese código, por lo que el parche no podía buscar el pedido de forma confiable.",
        51: "Añadí una columna y migración idempotente, y construí una prueba con dos usuarios. También comprobé que Bandit no detectaba la falta de autorización: era necesario probar el comportamiento, no solo escanear el código.",
        53: "Primero aislaría el endpoint y revisaría logs para saber qué pedidos y usuarios se afectaron; luego comunicaría el incidente según su alcance. No intentaría llamarlo resuelto solo porque el endpoint responde 202.",
        54: "Desplegaría la corrección con pruebas de propiedad y monitoreo de reenvíos. Además, configuraría un transporte de correo verificable o dejaría explícito que solo se registró un evento.",
    }
    for index, value in answers.items():
        replace_text(document.paragraphs[index], value)

    replace_text(document.tables[0].cell(4, 1).paragraphs[0], f"{REPO} (remediación {sha})")
    steps = document.tables[9]
    for row, explanation in {
        1: "Pendiente: reproducir el parche en la EC2 QA heredada; la integración y prueba roja actuales son locales.",
        2: "Pendiente: repetir en QA el bloqueo real. Localmente fallaron dos pruebas y la decisión fue BLOQUEADO.",
        3: "Clasificación local en docs/clasificacion_hallazgo.md; falta correlacionar con salida QA.",
        4: "Contención local y prevención separadas en docs/respuesta_incidente.md; comprobar contención QA.",
        5: "Autorización corregida y probada localmente; falta prueba de extremo a extremo en QA.",
        6: "Pendiente: repetir en QA la corrida verde local sobre el mismo commit corregido.",
        7: "Pendiente: crear la EC2 nueva solo tras QA verde y documentar su despliegue.",
    }.items():
        replace_text(steps.cell(row, 2).paragraphs[0], explanation)
    for row in (3, 4, 5):
        replace_text(steps.cell(row, 1).paragraphs[0], "[X]")
    for row in (1, 2, 6, 7):
        replace_text(steps.cell(row, 1).paragraphs[0], "[  ]")

    document.save(WORD)
    print(f"Actualizado: {WORD} (commit {sha})")


if __name__ == "__main__":
    if len(sys.argv) != 2 or len(sys.argv[1]) < 7:
        raise SystemExit("Especifica el SHA de remediación (mínimo 7 caracteres).")
    main(sys.argv[1])
